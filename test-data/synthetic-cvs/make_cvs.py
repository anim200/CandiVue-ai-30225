"""Make the 20 SYNTHETIC test CVs for the "PHP Laravel Developer" test job.

Every person here is made up (example.com emails, fake +880 1700-0000xx
phone numbers). Each CV says "SYNTHETIC TEST CV" at the top. Real CVs must
never be added to this folder (N13, N14, CN5).

The test job: minimum 5 years; mandatory PHP, Laravel and Docker;
preferred AWS certification. EXPECTED.md lists what screening should find.

Run on Windows (needs Microsoft Edge to print PDFs, python-docx for Word):
    python test-data/synthetic-cvs/make_cvs.py
"""

import html
import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).parent
BUILD = HERE / "_build"  # HTML and images used to make the PDFs (git-ignored)
EDGE = next(
    (
        p
        for p in (
            Path("C:/Program Files (x86)/Microsoft/Edge/Application/msedge.exe"),
            Path("C:/Program Files/Microsoft/Edge/Application/msedge.exe"),
        )
        if p.exists()
    ),
    None,
)
NOTE = "SYNTHETIC TEST CV: a made-up person for testing CandiVue AI. Not real personal data."


def job(title, company, dates, *bullets):
    """One work-history entry."""
    return {"title": title, "company": company, "dates": dates, "bullets": list(bullets)}


# Each CV: file name, layout and content. "headings" renames sections for the
# odd-layout edge case. Order: strong (1-5), partial (6-11), missing a
# mandatory requirement (12-16), edge cases (17-20).
CVS = [
    {
        "file": "01-jordan-sample", "layout": "classic",
        "name": "Jordan Sample", "role": "Senior PHP / Laravel Developer", "phone": "01", "city": "Dhaka",
        "summary": "PHP developer with 7 years of experience building Laravel web applications and REST APIs. Runs every project in Docker from development to production and deploys on AWS.",
        "jobs": [
            job("Senior PHP Developer", "Northwind Software Ltd., Dhaka", "Jan 2023 – Present",
                "Lead developer of a multi-tenant SaaS invoicing platform built with Laravel 11 and PHP 8.3.",
                "Containerised all services with Docker and Docker Compose.",
                "Deployed on AWS (EC2, RDS, S3, SQS) with GitHub Actions CI/CD.",
                "Introduced Redis queues and caching, cutting API response time by 45%."),
            job("PHP / Laravel Developer", "Example Tech Solutions, Dhaka", "Mar 2020 – Dec 2022",
                "Built REST APIs in Laravel 8/9 for an e-commerce order system.",
                "Wrote PHPUnit and Pest tests, raising coverage from 30% to 80%.",
                "Moved the legacy deployment to Docker images and Nginx + PHP-FPM containers."),
            job("Junior PHP Developer", "Contoso Web Studio, Chattogram", "Feb 2019 – Feb 2020",
                "Maintained client websites in PHP and Laravel 5.8 with MySQL."),
        ],
        "education": [("BSc in Computer Science and Engineering", "Sample University of Engineering and Technology", "2014 – 2018")],
        "certs": ["AWS Certified Developer – Associate (2024)", "Laravel Certified Developer (2021)"],
        "skills": {"Languages": "PHP 8.x, JavaScript, SQL", "Frameworks": "Laravel, Livewire, Vue.js",
                   "DevOps": "Docker, Docker Compose, GitHub Actions, Nginx", "Cloud": "AWS (EC2, RDS, S3, SQS)",
                   "Databases": "MySQL, PostgreSQL, Redis"},
        "projects": ["Invoicing SaaS: multi-tenant Laravel platform with Docker-based deployment on AWS."],
    },
    {
        "file": "02-alex-rowan", "layout": "modern",
        "name": "Alex Rowan", "role": "Lead Backend Engineer (PHP)", "phone": "03", "city": "Dhaka",
        "summary": "Backend engineer with 6 years in PHP and Laravel. Designs containerised microservices with Docker and Kubernetes on AWS. Holds two AWS certifications.",
        "jobs": [
            job("Lead Backend Engineer", "Fabrikam Payments, Dhaka", "Jun 2022 – Present",
                "Leads a team of 5 building payment APIs in Laravel 10 and PHP 8.2.",
                "Runs 14 Dockerised services on AWS EKS with Helm charts.",
                "Cut deployment time from 40 to 8 minutes with GitHub Actions and Docker layer caching."),
            job("Backend Developer", "Woodgrove Digital, Dhaka", "Aug 2020 – May 2022",
                "Built Laravel APIs for a ride-sharing back office; Docker Compose for local development.",
                "Wrote event-driven jobs with Laravel Horizon and Redis."),
            job("Software Engineer Intern", "Adventure Works BD, Dhaka", "Jan 2020 – Jul 2020",
                "Fixed bugs in a PHP 7 CRM and wrote MySQL reports."),
        ],
        "education": [("MSc in Software Engineering", "Example Institute of Technology", "2019 – 2021"),
                      ("BSc in Computer Science", "Sample University", "2014 – 2018")],
        "certs": ["AWS Certified Solutions Architect – Associate (2023)", "AWS Certified Developer – Associate (2022)"],
        "skills": {"Backend": "PHP 8, Laravel, Lumen, REST, GraphQL", "Infrastructure": "Docker, Kubernetes, Helm, Terraform, AWS",
                   "Data": "PostgreSQL, MySQL, Redis, Elasticsearch"},
        "projects": ["Payment gateway: PCI-DSS aware Laravel service handling 50,000 transactions per day."],
    },
    {
        "file": "03-sam-kader", "layout": "classic",
        "name": "Sam Kader", "role": "Senior Software Engineer", "phone": "04", "city": "Chattogram",
        "summary": "Software engineer with 8 years of professional PHP experience, the last 6 with Laravel. Comfortable owning a product end to end, from Docker-based local setups to cloud deployment on Google Cloud.",
        "jobs": [
            job("Senior Software Engineer", "Litware Commerce, Chattogram", "Apr 2021 – Present",
                "Owns the order and inventory modules of a Laravel 10 marketplace (1.2 million users).",
                "Introduced Docker for all environments and wrote the onboarding guide.",
                "Deploys to Google Cloud Run; manages Cloud SQL backups."),
            job("Software Engineer", "Proseware Ltd., Chattogram", "Jan 2019 – Mar 2021",
                "Migrated a plain-PHP shop to Laravel 6 step by step without downtime.",
                "Built a reporting module with queued Excel exports."),
            job("PHP Developer", "Coho Web Agency, Chattogram", "Feb 2018 – Dec 2018",
                "Built custom PHP and WordPress sites for small businesses."),
        ],
        "education": [("BSc in Computer Science and Engineering", "Sample Engineering University", "2011 – 2015")],
        "certs": ["Google Cloud Associate Cloud Engineer (2023)"],
        "skills": {"Languages": "PHP, JavaScript, Python", "Frameworks": "Laravel, Inertia.js, Vue",
                   "Tools": "Docker, Docker Compose, Git, Google Cloud, Redis, MySQL"},
        "projects": ["Marketplace search: Laravel Scout + Meilisearch, results in under 100 ms."],
    },
    {
        "file": "04-priya-example", "layout": "sidebar",
        "name": "Priya Example", "role": "PHP Laravel Developer", "phone": "05", "city": "Sylhet",
        "summary": "Laravel developer with 5 years 6 months of experience building SaaS dashboards and APIs. Every project I ship runs in Docker containers; I deploy and monitor on AWS.",
        "jobs": [
            job("PHP Laravel Developer", "Tailspin SaaS, Sylhet", "Feb 2022 – Present",
                "Built the subscription and billing module with Laravel Cashier.",
                "Wrote Dockerfiles and a Docker Compose setup used by the whole team.",
                "Set up AWS ECS Fargate deployments and CloudWatch alarms."),
            job("Web Developer", "Fourth Coffee Digital, Sylhet", "Apr 2021 – Jan 2022",
                "Developed Laravel admin panels and REST APIs for 12 clients.",
                "Used Docker for local development of every client project."),
        ],
        "education": [("BSc in Computer Science", "Sample Science and Technology University", "2015 – 2019")],
        "certs": ["AWS Certified Solutions Architect – Associate (2024)"],
        "skills": {"Skills": "PHP 8, Laravel, Docker, AWS ECS, MySQL, Redis, Vue.js, PHPUnit"},
        "projects": ["Billing module: recurring payments for 4,000 subscribers."],
    },
    {
        "file": "05-noor-testwell", "layout": "modern",
        "name": "Noor Testwell", "role": "Principal PHP Engineer", "phone": "06", "city": "Dhaka",
        "summary": "PHP engineer with 9 years of experience, Laravel since version 5. Architect of containerised platforms on AWS; also experienced with Symfony.",
        "jobs": [
            job("Principal PHP Engineer", "Blue Yonder Health, Dhaka", "Mar 2021 – Present",
                "Designed a HIPAA-style audit trail for a Laravel 10 health platform.",
                "Standardised Docker images and base layers for 9 PHP services on AWS.",
                "Mentors 8 engineers; runs architecture reviews."),
            job("Senior PHP Developer", "Datum Systems, Dhaka", "Jan 2018 – Feb 2021",
                "Built Symfony and Laravel services for logistics tracking; dockerised the stack."),
            job("PHP Developer", "Margie's Travel Tech, Dhaka", "Jul 2017 – Dec 2017",
                "Developed booking features in Laravel 5 and MySQL."),
        ],
        "education": [("BSc in Computer Science and Engineering", "Example University of Engineering", "2011 – 2015")],
        "certs": ["AWS Certified Solutions Architect – Professional (2022)", "Zend Certified PHP Engineer (2017)"],
        "skills": {"Core": "PHP 8.3, Laravel, Symfony", "Cloud and DevOps": "Docker, AWS (ECS, RDS, Lambda), Terraform",
                   "Practices": "DDD, TDD, code review"},
        "projects": ["Audit trail: immutable event log for 30 million records."],
    },
    {
        "file": "06-chris-mock", "layout": "classic",
        "name": "Chris Mock", "role": "PHP Developer", "phone": "07", "city": "Khulna",
        "summary": "PHP developer with 5 years of experience, mainly Laravel back-office applications. Uses Docker for local development.",
        "jobs": [
            job("PHP Developer", "Wingtip Toys Online, Khulna", "Oct 2023 – Present",
                "Maintains a Laravel 9 inventory system and its REST API.",
                "Runs the project locally with Docker Compose."),
            job("Junior PHP Developer", "Lucerne Publishing BD, Khulna", "Oct 2022 – Sep 2023",
                "Built Laravel forms, reports and email notifications."),
            job("Trainee Web Developer", "Lucerne Publishing BD, Khulna", "Oct 2021 – Sep 2022",
                "Wrote PHP and jQuery pages under supervision."),
        ],
        "education": [("Diploma in Computer Technology", "Sample Polytechnic Institute", "2015 – 2019")],
        "certs": [],
        "skills": {"Skills": "PHP, Laravel, MySQL, Docker (basic), Git, jQuery"},
        "projects": [],
    },
    {
        "file": "07-taylor-draft", "layout": "sidebar",
        "name": "Taylor Draft", "role": "Backend Developer", "phone": "08", "city": "Rajshahi",
        "summary": "Backend developer with 6 years of PHP experience; 3 years with Laravel. Familiar with Docker through our CI pipelines.",
        "jobs": [
            job("Backend Developer", "Northwind Traders BD, Rajshahi", "Jan 2023 – Present",
                "Develops Laravel 10 APIs for a logistics app.",
                "Maintains the GitLab CI pipeline that builds Docker images."),
            job("PHP Developer", "Graphic Design Institute, Rajshahi", "Sep 2021 – Dec 2022",
                "Built PHP 7 learning-management features without a framework.",
                "Moved two modules to Laravel 8."),
            job("Web Developer", "City Power BD, Rajshahi", "Oct 2020 – Aug 2021",
                "Maintained PHP intranet pages."),
        ],
        "education": [("BSc in Information and Communication Engineering", "Sample University", "2014 – 2018")],
        "certs": [],
        "skills": {"Skills": "PHP, Laravel, MySQL, GitLab CI, Docker, REST"},
        "projects": [],
    },
    {
        "file": "08-morgan-pilot", "layout": "modern",
        "name": "Morgan Pilot", "role": "Laravel Developer", "phone": "09", "city": "Dhaka",
        "summary": "Laravel developer with 5 years of experience. I set up every project with Laravel Sail, so my local and staging environments run in containers.",
        "jobs": [
            job("Laravel Developer", "Relecloud Apps, Dhaka", "Jun 2022 – Present",
                "Builds Laravel 10 apps with Livewire for a property-management startup.",
                "Uses Laravel Sail for all local environments and staging servers."),
            job("Junior Laravel Developer", "Trey Research BD, Dhaka", "Sep 2021 – May 2022",
                "Wrote Laravel controllers, Blade views and feature tests."),
        ],
        "education": [("BSc in Computer Science", "Example Private University", "2015 – 2019")],
        "certs": [],
        "skills": {"Stack": "PHP 8, Laravel, Livewire, Laravel Sail, MySQL, Tailwind CSS"},
        "projects": ["Rent collection: automated reminders and receipts for 900 tenants."],
    },
    {
        "file": "09-jamie-fixture", "layout": "classic",
        "name": "Jamie Fixture", "role": "PHP Developer", "phone": "10", "city": "Dhaka",
        "summary": "PHP developer with 7 years of experience, mostly WordPress and WooCommerce plugins; 2 years of Laravel. Uses Docker for local WordPress and Laravel environments.",
        "jobs": [
            job("PHP Developer", "Southridge Video, Dhaka", "Jan 2024 – Present",
                "Builds Laravel 11 APIs for a video subscription site.",
                "Uses Docker Compose for local environments."),
            job("WordPress Developer", "Alpine Ski House Web, Dhaka", "Mar 2020 – Dec 2023",
                "Built custom WordPress plugins and WooCommerce payment integrations in PHP."),
            job("Junior Web Developer", "Bellows College IT, Dhaka", "Nov 2019 – Feb 2020",
                "Built PHP pages for the college website."),
        ],
        "education": [("BSc in Computer Science", "Sample University", "2013 – 2017")],
        "certs": [],
        "skills": {"Skills": "PHP, WordPress, WooCommerce, Laravel, Docker, MySQL, JavaScript"},
        "projects": [],
    },
    {
        "file": "10-robin-stub", "layout": "sidebar",
        "name": "Robin Stub", "role": "PHP Microservices Developer", "phone": "11", "city": "Cumilla",
        "summary": "Developer with 5 years of experience building Laravel and Lumen microservices. I package services as OCI containers with Podman and run them on Kubernetes.",
        "jobs": [
            job("Backend Developer", "Humongous Insurance BD, Cumilla", "Aug 2022 – Present",
                "Builds Lumen microservices for claims processing.",
                "Packages services as OCI container images with Podman and Buildah."),
            job("Laravel Developer", "VanArsdel Retail, Cumilla", "Oct 2021 – Jul 2022",
                "Developed Laravel 7 store back-office features."),
        ],
        "education": [("BSc in Computer Science and Engineering", "Sample Technical University", "2015 – 2019")],
        "certs": ["Certified Kubernetes Application Developer (2023)"],
        "skills": {"Skills": "PHP, Laravel, Lumen, Podman, Kubernetes, PostgreSQL, RabbitMQ"},
        "projects": [],
    },
    {
        "file": "11-casey-proto", "layout": "classic",
        "name": "Casey Proto", "role": "Full-stack Laravel Developer", "phone": "12", "city": "Dhaka",
        "summary": "Full-stack developer with 6 years of Laravel and Vue.js experience, returning after a one-year career break (2023). Uses Docker for every project.",
        "jobs": [
            job("Full-stack Developer (contract)", "Contoso Pharma IT, Dhaka", "Feb 2024 – Present",
                "Builds Laravel 10 + Vue 3 dashboards for clinical-trial data.",
                "Maintains the Docker Compose development stack."),
            job("Laravel Developer", "Wide World Importers BD, Dhaka", "Jan 2019 – Dec 2022",
                "Built and maintained a Laravel import/export management system.",
                "Introduced Docker in 2020 for all environments."),
        ],
        "education": [("MSc in Computer Science", "Sample University", "2015 – 2016"),
                      ("BSc in Computer Science", "Sample University", "2011 – 2015")],
        "certs": [],
        "skills": {"Skills": "PHP, Laravel, Vue.js, Docker, MySQL, REST, Pest"},
        "projects": [],
    },
    {
        "file": "12-riley-example", "layout": "sidebar",
        "name": "Riley Example", "role": "PHP Web Developer", "phone": "02", "city": "Sylhet",
        "summary": "PHP developer with about 3 years of experience creating business websites, mostly with CodeIgniter and some Laravel. Interested in learning modern deployment tools.",
        "jobs": [
            job("PHP Developer", "Fabrikam Digital, Sylhet", "Aug 2024 – Present",
                "Builds company websites in CodeIgniter 4 with MySQL.",
                "Started maintaining a small Laravel 10 booking application.",
                "Deploys sites to shared hosting through FTP and cPanel."),
            job("Junior Web Developer", "Tailspin Web Agency, Sylhet", "Jan 2024 – Jul 2024",
                "Created responsive pages with Bootstrap and jQuery."),
        ],
        "education": [("Diploma in Computer Technology", "Sample Polytechnic Institute", "2018 – 2021")],
        "certs": [],
        "skills": {"Skills": "PHP, CodeIgniter, Laravel (basic), MySQL, jQuery, Bootstrap, Git"},
        "projects": [],
    },
    {
        "file": "13-dana-null", "layout": "modern",
        "name": "Dana Null", "role": "Senior PHP Developer", "phone": "13", "city": "Dhaka",
        "summary": "PHP developer with 8 years of experience in CodeIgniter and Yii2. Runs projects in Docker. Has not used Laravel professionally.",
        "jobs": [
            job("Senior PHP Developer", "Adatum ERP, Dhaka", "Jan 2020 – Present",
                "Leads development of a Yii2-based ERP; all services run in Docker.",
                "Designed the accounting module's database schema."),
            job("PHP Developer", "Consolidated Messenger BD, Dhaka", "Jan 2018 – Dec 2019",
                "Built SMS and email gateways in CodeIgniter 3."),
        ],
        "education": [("BSc in Computer Science and Engineering", "Sample University", "2011 – 2015")],
        "certs": [],
        "skills": {"Skills": "PHP, Yii2, CodeIgniter, Docker, MySQL, Redis"},
        "projects": [],
    },
    {
        "file": "14-lee-placeholder", "layout": "classic",
        "name": "Lee Placeholder", "role": "Laravel Developer", "phone": "14", "city": "Barishal",
        "summary": "Laravel developer with 6 years of experience building e-commerce and booking systems. Deploys to shared hosting and VPS servers with Laravel Forge.",
        "jobs": [
            job("Laravel Developer", "Tasmanian Traders BD, Barishal", "May 2022 – Present",
                "Builds Laravel 10 e-commerce features and payment integrations.",
                "Deploys with Laravel Forge to DigitalOcean servers."),
            job("PHP Developer", "Woodgrove Shop, Barishal", "Jun 2020 – Apr 2022",
                "Developed Laravel 5.6 booking forms and admin pages."),
        ],
        "education": [("BSc in Computer Science", "Sample University", "2014 – 2018")],
        "certs": [],
        "skills": {"Skills": "PHP, Laravel, MySQL, Laravel Forge, DigitalOcean, Vue.js"},
        "projects": [],
    },
    {
        "file": "15-kai-demo", "layout": "sidebar",
        "name": "Kai Demo", "role": "Junior Laravel Developer", "phone": "15", "city": "Dhaka",
        "summary": "Junior developer with 2 years of experience in Laravel and Docker. Fast learner who enjoys clean code.",
        "jobs": [
            job("Junior Laravel Developer", "Fineartschool Online, Dhaka", "Oct 2025 – Present",
                "Builds Laravel 11 features with Docker Compose for local development."),
            job("Intern, Web Development", "Northwind Software Ltd., Dhaka", "Oct 2024 – Sep 2025",
                "Wrote Laravel tests and fixed bugs."),
        ],
        "education": [("BSc in Computer Science and Engineering", "Sample University", "2019 – 2023")],
        "certs": [],
        "skills": {"Skills": "PHP, Laravel, Docker, MySQL, Git, Pest"},
        "projects": ["Student portal: final-year project in Laravel."],
    },
    {
        "file": "16-avery-mockup", "layout": "modern",
        "name": "Avery Mockup", "role": "Senior Node.js Developer", "phone": "16", "city": "Dhaka",
        "summary": "Backend developer with 6 years of experience in Node.js and TypeScript. Builds Dockerised services on AWS. No PHP experience.",
        "jobs": [
            job("Senior Node.js Developer", "Litware Streaming, Dhaka", "Jan 2022 – Present",
                "Builds NestJS microservices in Docker on AWS ECS."),
            job("Node.js Developer", "Proseware Cloud, Dhaka", "Jan 2020 – Dec 2021",
                "Built Express.js APIs and MongoDB data pipelines."),
        ],
        "education": [("BSc in Computer Science", "Sample University", "2014 – 2018")],
        "certs": ["AWS Certified Developer – Associate (2022)"],
        "skills": {"Skills": "Node.js, TypeScript, NestJS, Docker, AWS, MongoDB, PostgreSQL"},
        "projects": [],
    },
    {
        "file": "17-quinn-scan", "layout": "classic", "scanned": True,
        "name": "Quinn Scan", "role": "PHP Laravel Developer", "phone": "17", "city": "Dhaka",
        "summary": "Laravel developer with 6 years of experience. Docker in all environments; deploys on AWS. This CV is a scanned image (OCR test).",
        "jobs": [
            job("Laravel Developer", "Adventure Works BD, Dhaka", "Jan 2021 – Present",
                "Builds Laravel 10 REST APIs for a travel platform.",
                "Runs every service in Docker; deploys to AWS Elastic Beanstalk."),
            job("PHP Developer", "Coho Vineyard Tech, Dhaka", "Jan 2020 – Dec 2020",
                "Built Laravel 6 admin dashboards."),
        ],
        "education": [("BSc in Computer Science and Engineering", "Sample University", "2014 – 2018")],
        "certs": [],
        "skills": {"Skills": "PHP, Laravel, Docker, AWS, MySQL, Redis"},
        "projects": [],
    },
    {
        "file": "18-reese-docx", "layout": "docx",
        "name": "Reese Docx", "role": "PHP Laravel Developer", "phone": "18", "city": "Gazipur",
        "summary": "Laravel developer with 5 years of experience in API development. Uses Docker for development and testing. This CV is a Word file (DOCX test).",
        "jobs": [
            job("PHP Laravel Developer", "Lamna Healthcare BD, Gazipur", "Sep 2022 – Present",
                "Builds Laravel 10 APIs for appointment booking.",
                "Uses Docker Compose for local development and CI tests."),
            job("Junior PHP Developer", "Lamna Healthcare BD, Gazipur", "Sep 2021 – Aug 2022",
                "Built PHP and Laravel 7 internal tools."),
        ],
        "education": [("BSc in Computer Science", "Sample University", "2015 – 2019")],
        "certs": [],
        "skills": {"Skills": "PHP, Laravel, Docker, MySQL, PHPUnit"},
        "projects": [],
    },
    {
        "file": "19-skyler-twocol", "layout": "sidebar",
        "headings": {"summary": "About me", "jobs": "Where I've worked", "education": "Learning",
                     "certs": "Badges", "skills": "Toolbox", "projects": "Things I built"},
        "name": "Skyler Twocol", "role": "Laravel Engineer", "phone": "19", "city": "Dhaka",
        "summary": "Laravel engineer, about seven years in the PHP world. Everything I build ships in Docker containers to AWS.",
        "jobs": [
            job("Laravel Engineer", "Tailwind Traders BD, Dhaka", "03/2022 – now",
                "Laravel 11 back office for a retail chain; Dockerised and on AWS ECS."),
            job("PHP developer", "Fabrikam Residences, Dhaka", "11/2018 – 02/2022",
                "Laravel 6–9 property portal; moved it into Docker in 2020."),
        ],
        "education": [("B.Sc. (Engg.) CSE", "Sample University", "2014 – 2018")],
        "certs": ["AWS Certified Cloud Practitioner (2023)"],
        "skills": {"Toolbox": "PHP 8 · Laravel · Docker · AWS ECS · MySQL · Redis · Vue"},
        "projects": ["Retail back office: stock transfers across 40 shops."],
    },
    {
        "file": "20-drew-minimal", "layout": "minimal",
        "name": "Drew Minimal", "role": "PHP Developer", "phone": "20", "city": "Dhaka",
        "summary": "",
        "jobs": [
            job("PHP Developer", "Self-employed / freelance", "2016 – present",
                "Laravel and Docker projects for clients in Europe."),
        ],
        "education": [],
        "certs": [],
        "skills": {"Skills": "PHP, Laravel, Docker"},
        "projects": [],
    },
]

BASE_CSS = """
@page { size: A4; margin: 0; }
body { margin: 0; color: #1a2227; font-size: 10.5pt; line-height: 1.45; }
.note { font: 8pt Arial, sans-serif; color: #8a2d2d; border: 1px solid #e5b4b4;
        background: #fdf2f2; padding: 4px 8px; margin-bottom: 12px; }
ul { margin: 4px 0 8px; padding-left: 18px; } li { margin: 2px 0; }
.job { margin-bottom: 10px; } .head { font-weight: 600; }
.sub { color: #555; font-size: 9.5pt; }
"""

LAYOUT_CSS = {
    "classic": """body { font-family: "Segoe UI", Arial, sans-serif; padding: 16mm; }
        h1 { font-size: 22pt; margin: 0; color: #103a40; } .role { color: #23717a; font-size: 12pt; }
        h2 { font-size: 11pt; text-transform: uppercase; letter-spacing: .06em; color: #23717a;
             border-bottom: 1.5px solid #d5e9eb; margin: 14px 0 6px; padding-bottom: 2px; }""",
    "modern": """body { font-family: Calibri, Arial, sans-serif; }
        .band { background: #2b3a55; color: #fff; padding: 12mm 16mm 8mm; }
        .band h1 { margin: 0; font-size: 24pt; } .band .role { color: #c9d6ff; }
        .content { padding: 6mm 16mm 12mm; }
        h2 { font-size: 12pt; color: #2b3a55; margin: 14px 0 6px; }
        h2::after { content: ""; display: block; width: 40px; height: 3px; background: #e0a526; margin-top: 3px; }""",
    "sidebar": """body { font-family: Georgia, serif; }
        .page { display: grid; grid-template-columns: 62mm 1fr; min-height: 297mm; }
        .side { background: #2f3e46; color: #eef2f3; padding: 16mm 8mm 12mm 10mm; font-size: 9.5pt; }
        .side h2 { color: #cad2c5; font-size: 10pt; text-transform: uppercase; letter-spacing: .08em; }
        .main { padding: 16mm 14mm 12mm 10mm; }
        h1 { font-size: 20pt; margin: 0; color: #2f3e46; } .role { color: #52796f; }
        h2 { font-size: 12pt; color: #2f3e46; border-bottom: 1px solid #cad2c5; margin: 14px 0 6px; }""",
    "minimal": """body { font-family: "Courier New", monospace; padding: 20mm; font-size: 10pt; }
        h1 { font-size: 16pt; margin: 0; } h2 { font-size: 10pt; margin: 14px 0 4px; }""",
}


def esc(text: str) -> str:
    """Escape text for HTML."""
    return html.escape(text)


def sections(cv: dict) -> dict:
    """The HTML of each section (empty sections are left out)."""
    names = {"summary": "Professional Summary", "jobs": "Work Experience", "education": "Education",
             "certs": "Certifications", "skills": "Skills", "projects": "Projects"}
    names.update(cv.get("headings", {}))
    parts = {}
    if cv["summary"]:
        parts["summary"] = f"<h2>{esc(names['summary'])}</h2><p>{esc(cv['summary'])}</p>"
    jobs = "".join(
        f'<div class="job"><div class="head">{esc(j["title"])} · {esc(j["dates"])}</div>'
        f'<div class="sub">{esc(j["company"])}</div>'
        f'<ul>{"".join(f"<li>{esc(b)}</li>" for b in j["bullets"])}</ul></div>'
        for j in cv["jobs"]
    )
    parts["jobs"] = f"<h2>{esc(names['jobs'])}</h2>{jobs}"
    if cv["education"]:
        parts["education"] = f"<h2>{esc(names['education'])}</h2>" + "".join(
            f'<div class="job"><div class="head">{esc(d)}</div><div class="sub">{esc(s)} · {esc(y)}</div></div>'
            for d, s, y in cv["education"]
        )
    if cv["certs"]:
        parts["certs"] = f"<h2>{esc(names['certs'])}</h2><ul>" + "".join(
            f"<li>{esc(c)}</li>" for c in cv["certs"]
        ) + "</ul>"
    parts["skills"] = f"<h2>{esc(names['skills'])}</h2>" + "".join(
        f"<div><b>{esc(k)}:</b> {esc(v)}</div>" for k, v in cv["skills"].items()
    )
    if cv["projects"]:
        parts["projects"] = f"<h2>{esc(names['projects'])}</h2><ul>" + "".join(
            f"<li>{esc(p)}</li>" for p in cv["projects"]
        ) + "</ul>"
    return parts


def contact(cv: dict) -> str:
    """Email, phone and city (all made up)."""
    email = cv["name"].lower().replace(" ", ".") + "@example.com"
    return f"{esc(cv['city'])}, Bangladesh · {email} · +880 1700-0000{cv['phone']}"


def page_html(cv: dict) -> str:
    """The whole CV as an HTML page in its layout."""
    s = sections(cv)
    head = f"<h1>{esc(cv['name'])}</h1><div class='role'>{esc(cv['role'])}</div>"
    if cv["layout"] == "sidebar":
        side = f"<h2>Contact</h2><p>{contact(cv).replace(' · ', '<br>')}</p>" + s.pop("skills") + s.pop("certs", "")
        body = (f"<div class='page'><aside class='side'>{side}</aside><main class='main'>"
                f"<div class='note'>{NOTE}</div>{head}{''.join(s.values())}</main></div>")
    elif cv["layout"] == "modern":
        body = (f"<div class='band'><div class='note' style='color:#8a2d2d'>{NOTE}</div>{head}"
                f"<div>{contact(cv)}</div></div><div class='content'>{''.join(s.values())}</div>")
    else:
        body = f"<div class='note'>{NOTE}</div>{head}<div class='sub'>{contact(cv)}</div>{''.join(s.values())}"
    css = BASE_CSS + LAYOUT_CSS["classic" if cv["layout"] == "docx" else cv["layout"]]
    return f"<!doctype html><html><head><meta charset='utf-8'><style>{css}</style></head><body>{body}</body></html>"


def edge(*args: str) -> None:
    """Run headless Microsoft Edge (prints PDFs and screenshots)."""
    profile = BUILD / "edge-profile"
    subprocess.run([str(EDGE), "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                    f"--user-data-dir={profile}", *args], check=True, capture_output=True, timeout=120)


def make_pdf(cv: dict) -> None:
    """Print the CV's HTML to PDF."""
    source = BUILD / f"{cv['file']}.html"
    source.write_text(page_html(cv), encoding="utf-8")
    edge(f"--print-to-pdf={HERE / (cv['file'] + '.pdf')}", source.as_uri())


def make_scanned_pdf(cv: dict) -> None:
    """A PDF that holds only a picture of the CV (no text layer), like a scan."""
    source = BUILD / f"{cv['file']}.html"
    source.write_text(page_html(cv), encoding="utf-8")
    image = BUILD / f"{cv['file']}.png"
    edge("--hide-scrollbars", "--window-size=1240,1754", f"--screenshot={image}", source.as_uri())
    scan = BUILD / f"{cv['file']}-scan.html"
    scan.write_text(
        "<!doctype html><html><head><style>@page{size:A4;margin:0}body{margin:0;background:#f4f1ea}"
        "img{width:210mm;filter:grayscale(1) contrast(1.15) blur(0.35px);transform:rotate(0.6deg)}"
        f"</style></head><body><img src='{image.name}'></body></html>",
        encoding="utf-8",
    )
    edge(f"--print-to-pdf={HERE / (cv['file'] + '.pdf')}", scan.as_uri())


def make_docx(cv: dict) -> None:
    """Write the CV as a Word document."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_paragraph(NOTE)
    doc.add_heading(cv["name"], level=0)
    doc.add_paragraph(f"{cv['role']}\n{contact(cv)}")
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(cv["summary"])
    doc.add_heading("Work Experience", level=1)
    for j in cv["jobs"]:
        doc.add_paragraph(f"{j['title']} · {j['dates']}", style="Heading 3")
        doc.add_paragraph(j["company"])
        for b in j["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
    doc.add_heading("Education", level=1)
    for d, s, y in cv["education"]:
        doc.add_paragraph(f"{d}, {s} ({y})")
    doc.add_heading("Skills", level=1)
    table = doc.add_table(rows=0, cols=2)
    for k, v in cv["skills"].items():
        row = table.add_row().cells
        row[0].text, row[1].text = k, v
    doc.save(HERE / f"{cv['file']}.docx")


def main() -> None:
    """Make all 20 CVs in this folder."""
    if EDGE is None:
        raise SystemExit("Microsoft Edge was not found: it is needed to print the PDFs.")
    BUILD.mkdir(exist_ok=True)
    for cv in CVS:
        if cv["layout"] == "docx":
            make_docx(cv)
        elif cv.get("scanned"):
            make_scanned_pdf(cv)
        else:
            make_pdf(cv)
        print("made", cv["file"])
    shutil.rmtree(BUILD / "edge-profile", ignore_errors=True)


if __name__ == "__main__":
    main()
