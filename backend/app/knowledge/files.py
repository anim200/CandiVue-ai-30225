"""Checks an uploaded document and reads its text (N8).

The type is decided by the file's content, not by its name: a file called
"policy.pdf" that is really something else is refused. Only the text is
kept; the file itself is never stored (N16).
"""

import io
import re
import zipfile

from docx import Document
from pypdf import PdfReader

# Messages the user sees (plain language, no technical details, N12).
UNSUPPORTED = "Only PDF, DOCX, TXT or MD files can be added."
UNREADABLE = "This file could not be read. It may be damaged or password-protected."
NO_TEXT = (
    "No text was found in this file. Scanned documents need text to be searchable."
)

TEXT_EXTENSIONS = {".txt": "txt", ".md": "md"}


class DocumentError(Exception):
    """The upload cannot be used; the message is safe to show to the user."""


def clean_filename(filename: str) -> str:
    """Keep only a safe, short display name (no folders, no control characters)."""
    name = re.split(r"[\\/]", filename or "")[-1]  # drop any folder part
    name = re.sub(r"[\x00-\x1f\x7f]", "", name).strip()
    return name[:255] or "document"


def detect_file_type(data: bytes, filename: str) -> str:
    """Return "pdf", "docx", "txt" or "md" from the file's content, or raise."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):  # a zip file: is it a Word document?
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
        raise DocumentError(UNSUPPORTED)
    # Plain text: allowed only with a .txt/.md name and valid UTF-8 text.
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension in TEXT_EXTENSIONS and b"\x00" not in data:
        try:
            data.decode("utf-8")
            return TEXT_EXTENSIONS[extension]
        except UnicodeDecodeError:
            pass
    raise DocumentError(UNSUPPORTED)


def extract_text(data: bytes, file_type: str) -> str:
    """Read all text from the document; raise DocumentError if there is none."""
    try:
        if file_type == "pdf":
            pages = PdfReader(io.BytesIO(data)).pages
            text = "\n\n".join(page.extract_text() or "" for page in pages)
        elif file_type == "docx":
            document = Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:  # text inside tables too
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(parts)
        else:
            text = data.decode("utf-8")
    except DocumentError:
        raise
    except Exception as exc:  # damaged or encrypted files
        raise DocumentError(UNREADABLE) from exc

    if not text.strip():
        raise DocumentError(NO_TEXT)
    return text
